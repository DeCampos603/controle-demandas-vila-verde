"""Gera a planilha limpa do Controle de Demandas (mesmo layout que Codigo.gs > configurar()).

Uso: py -3 ferramentas/gerar_planilha.py
Saída: Demandas-Vila-Verde.xlsx na raiz do projeto. Abra no Google Sheets e use
Arquivo > Salvar como Planilha Google; é essa cópia nativa que o Apps Script usa.
Os cabeçalhos DEVEM ser idênticos aos de CAMPOS em apps-script/Codigo.gs.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUNAS = [  # (cabeçalho, largura, lista opcional)
    ("ID do Protocolo", 16, None),
    ("Data de Abertura", 18, None),
    ("Última Atualização", 18, None),
    ("Data de Conclusão", 18, None),
    ("Categoria", 22, ["Adequação", "Reforma", "Melhoria", "Manutenção Corretiva"]),
    ("Local / PNR", 22, None),
    ("Descrição do Escopo", 55, None),
    ("Prioridade", 26, ["Nível 1 (Urgente/Crítico)", "Nível 2 (Moderado)", "Nível 3 (Rotina/Baixa)"]),
    ("Matriz de Risco", 16, ["Estrutural", "Elétrico", "Sanitário", "Nenhum"]),
    ("Benefício Direto", 40, None),
    ("Despacho Institucional", 40, None),
    ("Status da Execução", 18, ["A começar", "Em andamento", "Pendência", "Finalizado"]),
    ("Motivo da Pendência", 32, None),
    ("Excluído", 11, ["SIM"]),
]
DATAS = {2, 3, 4}
LINHAS = 2000
VERDE, MARINHO = "1E6B34", "1B2A4A"


def main() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Demandas"
    for i, (nome, larg, lista) in enumerate(COLUNAS, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = larg
        letra = c.column_letter
        for r in range(2, LINHAS + 1):
            cel = ws.cell(row=r, column=i)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.number_format = "dd/mm/yyyy hh:mm" if i in DATAS else "@"
        if lista:
            dv = DataValidation(type="list", formula1='"' + ",".join(lista) + '"', allow_blank=True, showErrorMessage=True,
                                errorTitle="Valor inválido", error="Escolha um item da lista.")
            dv.add(f"{letra}2:{letra}{LINHAS}")
            ws.add_data_validation(dv)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUNAS)).column_letter}1"

    h = wb.create_sheet("Historico")
    for i, (nome, larg) in enumerate([("Quando", 20), ("ID do Protocolo", 16), ("Ação", 12), ("Campo", 24), ("De", 40), ("Para", 40)], start=1):
        c = h.cell(row=1, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MARINHO)
        h.column_dimensions[c.column_letter].width = larg
    h.freeze_panes = "A2"

    destino = Path(__file__).resolve().parent.parent / "Demandas-Vila-Verde.xlsx"
    wb.save(destino)
    return destino


if __name__ == "__main__":
    print(main())
